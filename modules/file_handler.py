import os
import PyPDF2
import docx
import openpyxl

def read_pdf(file_path):
    text = ""
    with open(file_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text += page.extract_text() + "\n"
    return text

def read_docx(file_path):
    doc = docx.Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])

def read_txt(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def read_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path)
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) if cell else "" for cell in row])
            text += row_text + "\n"
    return text

def load_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return read_pdf(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    elif ext == ".txt":
        return read_txt(file_path)
    elif ext == ".xlsx":
        return read_xlsx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
